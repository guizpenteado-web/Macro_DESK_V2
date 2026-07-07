"""
Gerador do relatorio Excel.
Planilha com 3 abas: Resumo, Breadth Historico, Ativos.
"""
from __future__ import annotations
from datetime import datetime
from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows
from app.config.settings import settings
from app.database.connection import engine
from app.utils.logger import logger


def _pct_fill(pct: float) -> PatternFill:
    if pct >= 70: color = "2ea043"
    elif pct >= 50: color = "e3b341"
    elif pct >= 30: color = "d29922"
    else: color = "f85149"
    return PatternFill("solid", fgColor=color)


def generate_report() -> Path:
    breadth = pd.read_sql("SELECT * FROM breadth ORDER BY date DESC LIMIT 504", engine)
    indicators = pd.read_sql(
        """SELECT i.ticker, a.name, a.weight, i.date, i.close,
                  i.sma21, i.sma50, i.sma200,
                  i.above_sma21, i.above_sma50, i.above_sma200
           FROM indicators i LEFT JOIN assets a ON a.ticker=i.ticker
           WHERE i.date=(SELECT MAX(date) FROM indicators WHERE ticker=i.ticker)
           ORDER BY a.weight DESC NULLS LAST, i.ticker""",
        engine,
    )

    wb = Workbook()

    # ── Aba 1: Resumo ────────────────────────────────────────────
    ws = wb.active
    ws.title = "Resumo"
    ws.sheet_view.showGridLines = False
    ws["A1"] = "IBOV Market Breadth"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Gerado em: {datetime.now().strftime("%d/%m/%Y %H:%M")}"
    ws["A2"].font = Font(color="888888", size=10)

    if not breadth.empty:
        latest = breadth.iloc[0]
        ws["A4"] = "Indicador"
        ws["B4"] = "Valor"
        ws["C4"] = "Ativos"
        ws["A4"].font = ws["B4"].font = ws["C4"].font = Font(bold=True)
        for i, (label, pct_col, cnt_col) in enumerate([
            ("SMA21",  "pct_sma21",  "above_sma21"),
            ("SMA50",  "pct_sma50",  "above_sma50"),
            ("SMA200", "pct_sma200", "above_sma200"),
        ], start=5):
            ws[f"A{i}"] = f"% Acima {label}"
            ws[f"B{i}"] = round(float(latest[pct_col]), 1)
            ws[f"B{i}"].number_format = "0.0"
            ws[f"B{i}"].fill = _pct_fill(float(latest[pct_col]))
            ws[f"B{i}"].font = Font(bold=True, color="FFFFFF")
            ws[f"C{i}"] = str(int(latest[cnt_col])) + " de " + str(int(latest["total_assets"]))

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 18

    # ── Aba 2: Breadth Historico ─────────────────────────────────
    ws2 = wb.create_sheet("Breadth Historico")
    breadth_fmt = breadth.rename(columns={
        "date":"Data","total_assets":"Total","above_sma21":"Acima SMA21",
        "above_sma50":"Acima SMA50","above_sma200":"Acima SMA200",
        "pct_sma21":"% SMA21","pct_sma50":"% SMA50","pct_sma200":"% SMA200"
    })
    for r_idx, row in enumerate(dataframe_to_rows(breadth_fmt, index=False, header=True), 1):
        ws2.append(row)
        if r_idx == 1:
            for cell in ws2[1]:
                cell.font = Font(bold=True)

    # ── Aba 3: Ativos ────────────────────────────────────────────
    ws3 = wb.create_sheet("Ativos")
    for r_idx, row in enumerate(dataframe_to_rows(indicators, index=False, header=True), 1):
        ws3.append(row)
        if r_idx == 1:
            for cell in ws3[1]:
                cell.font = Font(bold=True)
        if r_idx > 1:
            row_data = indicators.iloc[r_idx - 2]
            for col, flag_col in [(10,"above_sma21"),(11,"above_sma50"),(12,"above_sma200")]:
                cell = ws3.cell(row=r_idx, column=col)
                if row_data[flag_col]:
                    cell.fill = PatternFill("solid", fgColor="2ea04330")

    for ws_obj in [ws2, ws3]:
        ws_obj.sheet_view.showGridLines = False
        for col in ws_obj.columns:
            max_len = max((len(str(c.value or "")) for c in col), default=8)
            ws_obj.column_dimensions[col[0].column_letter].width = min(max_len + 2, 30)

    out = settings.reports_output / "ibov_breadth_report.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    logger.success(f"Relatorio Excel gerado: {out}")
    return out